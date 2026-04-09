"""Convert PDF/DOCX/PPTX files to Markdown for qmd indexing.

Usage:
    py scripts/binary_docs_to_markdown.py
    py scripts/binary_docs_to_markdown.py --input "docs/医院材料学习" --output "docs/knowledge-base/.staging-binary-md"
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import fitz  # pymupdf
from docx import Document
from pptx import Presentation


def safe_stem(name: str) -> str:
    base = Path(name).stem
    base = re.sub(r"\s+", "-", base.strip())
    base = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", base)
    base = re.sub(r"-{2,}", "-", base).strip("-")
    return base or "untitled"


def pdf_to_md(path: Path) -> str:
    chunks: list[str] = [f"# {path.stem}", "", f"**Source:** `{path.name}`", ""]
    with fitz.open(path) as doc:
        for idx, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if not text:
                continue
            chunks.extend([f"## Page {idx}", "", text, ""])
    return "\n".join(chunks).strip() + "\n"


def docx_to_md(path: Path) -> str:
    doc = Document(path)
    chunks: list[str] = [f"# {path.stem}", "", f"**Source:** `{path.name}`", ""]
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            chunks.append(text)
    return "\n\n".join(chunks).strip() + "\n"


def pptx_to_md(path: Path) -> str:
    prs = Presentation(path)
    chunks: list[str] = [f"# {path.stem}", "", f"**Source:** `{path.name}`", ""]
    for idx, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                t = shape.text.strip()
                if t:
                    parts.append(t)
        if parts:
            chunks.extend([f"## Slide {idx}", "", "\n\n".join(parts), ""])
    return "\n".join(chunks).strip() + "\n"


def convert_file(src: Path, out_dir: Path) -> Path | None:
    ext = src.suffix.lower()
    if ext not in {".pdf", ".docx", ".pptx"}:
        return None
    if ext == ".pdf":
        md = pdf_to_md(src)
    elif ext == ".docx":
        md = docx_to_md(src)
    else:
        md = pptx_to_md(src)
    if not md.strip():
        return None
    out = out_dir / f"{safe_stem(src.name)}.md"
    out.write_text(md, encoding="utf-8")
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert binary docs to markdown")
    parser.add_argument("--input", default=r"docs\医院材料学习")
    parser.add_argument("--output", default=r"docs\knowledge-base\.staging-binary-md")
    args = parser.parse_args()

    input_dir = Path(args.input)
    out_dir = Path(args.output)
    out_dir.mkdir(parents=True, exist_ok=True)

    files = sorted(
        [
            p
            for p in input_dir.iterdir()
            if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".pptx"}
        ],
        key=lambda p: p.name.lower(),
    )
    print(f"Found {len(files)} binary document(s) in {input_dir}")
    created = 0
    for f in files:
        out = convert_file(f, out_dir)
        if out:
            created += 1
            print(f"  Converted: {f.name} -> {out.name}")
    print(f"Done. Created {created} markdown file(s) in {out_dir}")


if __name__ == "__main__":
    main()

