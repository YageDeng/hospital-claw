"""Convert XLSX files to Markdown tables for MinerU ingestion.

Usage:
    python scripts/xlsx_to_markdown.py
    python scripts/xlsx_to_markdown.py --input "docs/医院材料学习" --output "docs/knowledge-base/.staging"
"""

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def sheet_to_markdown(ws) -> str:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ""

    headers = [str(cell) if cell is not None else "" for cell in rows[0]]
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows[1:]:
        cells = [str(cell) if cell is not None else "" for cell in row]
        if all(c == "" for c in cells):
            continue
        lines.append("| " + " | ".join(cells) + " |")

    return "\n".join(lines)


def convert_xlsx(xlsx_path: Path, output_dir: Path) -> list[Path]:
    wb = load_workbook(xlsx_path, data_only=True)
    created = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        md_table = sheet_to_markdown(ws)
        if not md_table:
            continue

        safe_name = xlsx_path.stem
        if len(wb.sheetnames) > 1:
            safe_name += f"_{sheet_name}"

        output_path = output_dir / f"{safe_name}.md"
        content = (
            f"# {xlsx_path.stem}\n\n"
            f"**Sheet:** {sheet_name}\n\n"
            f"**Source:** `{xlsx_path.name}`\n\n"
            f"{md_table}\n"
        )
        output_path.write_text(content, encoding="utf-8")
        created.append(output_path)
        print(f"  Converted: {xlsx_path.name} / {sheet_name} -> {output_path.name}")

    return created


def main():
    parser = argparse.ArgumentParser(description="Convert XLSX to Markdown for MinerU")
    parser.add_argument(
        "--input",
        default=r"docs\医院材料学习",
        help="Directory containing XLSX files",
    )
    parser.add_argument(
        "--output",
        default=r"docs\knowledge-base\.staging",
        help="Output directory for markdown files",
    )
    args = parser.parse_args()

    input_dir = Path(args.input)
    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)

    xlsx_files = list(input_dir.glob("*.xlsx"))
    if not xlsx_files:
        print(f"No XLSX files found in {input_dir}")
        return

    print(f"Found {len(xlsx_files)} XLSX file(s) in {input_dir}")
    print(f"Output directory: {output_dir}\n")

    all_created = []
    for xlsx_path in xlsx_files:
        print(f"Processing: {xlsx_path.name}")
        created = convert_xlsx(xlsx_path, output_dir)
        all_created.extend(created)

    print(f"\nDone. Created {len(all_created)} markdown file(s).")


if __name__ == "__main__":
    main()